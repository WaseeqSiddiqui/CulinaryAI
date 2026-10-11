from __future__ import annotations

from datetime import date
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "CulinaryAI_Project_Progress_and_Roadmap.docx"

ORANGE = "E85D35"
GREEN = "16856B"
DARK = "20242A"
MUTED = "667085"
LIGHT = "F6F2EC"
PALE_ORANGE = "FFF1EB"
PALE_GREEN = "EAF7F2"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def add_toc_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Update this field in Microsoft Word to refresh the table of contents."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size, color in (
        ("Title", 34, ORANGE),
        ("Subtitle", 15, MUTED),
        ("Heading 1", 22, ORANGE),
        ("Heading 2", 15, DARK),
        ("Heading 3", 11.5, GREEN),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = style_name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if style_name != "Heading 1" else 18)
        style.paragraph_format.space_after = Pt(5)

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Block"]
    code_style.font.name = "Consolas"
    code_style.font.size = Pt(8.5)
    code_style.font.color.rgb = RGBColor.from_string(DARK)
    code_style.paragraph_format.left_indent = Inches(0.18)
    code_style.paragraph_format.right_indent = Inches(0.18)
    code_style.paragraph_format.space_before = Pt(4)
    code_style.paragraph_format.space_after = Pt(7)

    if "Callout" not in styles:
        callout_style = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
    else:
        callout_style = styles["Callout"]
    callout_style.font.name = "Aptos"
    callout_style.font.size = Pt(10.5)
    callout_style.font.bold = True
    callout_style.font.color.rgb = RGBColor.from_string(GREEN)
    callout_style.paragraph_format.left_indent = Inches(0.18)
    callout_style.paragraph_format.right_indent = Inches(0.18)
    callout_style.paragraph_format.space_before = Pt(5)
    callout_style.paragraph_format.space_after = Pt(8)

    for section in document.sections:
        header = section.header.paragraphs[0]
        header.text = "CULINARYAI  /  PROJECT PROGRESS & ROADMAP"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.runs[0].font.name = "Aptos"
        header.runs[0].font.size = Pt(8)
        header.runs[0].font.bold = True
        header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
        add_page_number(section.footer.paragraphs[0])

    props = document.core_properties
    props.title = "CulinaryAI Project Progress and Future Roadmap"
    props.subject = "Detailed implementation report for the CulinaryAI AI/ML course project"
    props.author = "CulinaryAI Project Team"
    props.keywords = "Flutter, Firebase, FastAPI, RAG, AI, ML, nutrition, pantry"


def add_rule(document: Document, color: str = ORANGE) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "14")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def add_callout(document: Document, text: str, fill: str = PALE_GREEN) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 130, 170, 130, 170)
    paragraph = cell.paragraphs[0]
    paragraph.style = document.styles["Callout"]
    paragraph.add_run(text)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code_flow(document: Document, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT)
    set_cell_margins(cell, 140, 160, 140, 160)
    paragraph = cell.paragraphs[0]
    paragraph.style = document.styles["Code Block"]
    paragraph.add_run(text)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullets(document: Document, items: list[str], level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    for item in items:
        paragraph = document.add_paragraph(item, style=style)
        paragraph.paragraph_format.space_after = Pt(3)


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(item, style="List Number")
        paragraph.paragraph_format.space_after = Pt(4)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths=None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, heading in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, ORANGE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(heading)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
        run.font.size = Pt(9)
        if widths:
            cell.width = Inches(widths[index])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if row_index % 2 == 1:
                set_cell_shading(cell, "FAFAFA")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(str(value))
            run.font.size = Pt(8.6)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_cover(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(70)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    icon = paragraph.add_run("CULINARYAI")
    icon.bold = True
    icon.font.name = "Aptos Display"
    icon.font.size = Pt(18)
    icon.font.color.rgb = RGBColor.from_string(GREEN)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Project Progress\nand Future Roadmap")

    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("AI-Powered Household Meal Planning & Kitchen Management System")

    add_rule(document)

    meta = document.add_table(rows=5, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.style = "Table Grid"
    details = [
        ("Document purpose", "Deep implementation status, technical explanation, and planned development"),
        ("Project type", "AI/ML course project — Flutter mobile app with Firebase and FastAPI"),
        ("Current milestone", "Week 1 implementation completed and verified"),
        ("Repository branch", "main"),
        ("Prepared", date.today().strftime("%d %B %Y")),
    ]
    for row, (label, value) in zip(meta.rows, details):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], DARK)
        row.cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(WHITE)
        row.cells[0].paragraphs[0].runs[0].bold = True
        for cell in row.cells:
            set_cell_margins(cell, 120, 140, 120, 140)

    document.add_paragraph()
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run(
        "Status is based on the repository at commit c21f739 and verified test results."
    )
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    document.add_page_break()


def build_report() -> Document:
    document = Document()
    configure_document(document)
    add_cover(document)

    document.add_heading("Table of Contents", level=1)
    toc = document.add_paragraph()
    add_toc_field(toc)
    document.add_paragraph(
        "If the page numbers do not appear immediately, right-click the table in Microsoft Word and select Update Field."
    )
    document.add_page_break()

    document.add_heading("1. Executive Summary", level=1)
    document.add_paragraph(
        "CulinaryAI is an AI-assisted kitchen management application designed to connect household inventory, recipe discovery, guided cooking, nutrition, receipt ingestion, and replenishment in one workflow. The mobile client is written in Flutter, identity and shared data are handled through Firebase, and AI-oriented meal-planning services are separated into a Python FastAPI microservice."
    )
    document.add_paragraph(
        "The Week 1 milestone has produced a working application foundation rather than a collection of disconnected mock screens. Authentication persists between sessions; households can be created or joined; pantry items can be stored, filtered, updated, shared, and restored offline; receipt captures flow through a review screen; nutrition and portion calculations are deterministic and tested; recipes can launch guided cooking; a domain chatbot can answer pantry and recipe questions; and a FastAPI service can retrieve verified recipes, rank pantry matches, scale ingredients, and return structured cooking steps."
    )
    add_callout(
        document,
        "Important maturity statement: the current chatbot is a local culinary question-answering layer, receipt extraction uses a Week 1 reviewable fixture, and backend retrieval uses deterministic hash embeddings. These are functional prototypes with stable interfaces; real OCR, neural embeddings, vector search, and learned personalization are future upgrades.",
        PALE_ORANGE,
    )

    add_table(
        document,
        ["Area", "Current status", "Evidence / outcome"],
        [
            ["Flutter application shell", "Completed", "Five-tab navigation, responsive dashboard, drawer, light/dark theme"],
            ["Authentication and sessions", "Completed", "Firebase-ready sign-in/sign-up plus persistent local session fallback"],
            ["Shared household pantry", "Completed", "Create/join household, scoped pantry data, offline cache and Firestore rules"],
            ["Receipt workflow", "Prototype completed", "Camera/gallery capture, processing states, editable confirmation, pantry batch add; OCR model pending"],
            ["Nutrition and scaling", "Completed", "Unit normalization, macro totals, per-serving values, ingredient scaling"],
            ["Recipe and guided cooking", "Completed for local catalog", "Pantry match, missing ingredients, detail view, cooking steps, completion deduction"],
            ["Chef Chat", "Completed for domain intents", "Answers pantry, stock, recipe, time, calorie and instruction questions"],
            ["FastAPI / RAG foundation", "Week 1 completed", "Verified recipe repository, local embeddings, retrieval, /generate-plan and instruction endpoint"],
            ["Production AI/ML", "Planned", "Semantic embedding model, vector database, ranker, feedback learning, real OCR"],
        ],
        [1.45, 1.25, 4.2],
    )

    document.add_heading("2. Project Objectives and Scope", level=1)
    document.add_paragraph(
        "The core problem is that households often decide meals without a reliable view of available ingredients, allow food to expire, duplicate purchases, and switch between unrelated applications for recipes, nutrition, inventory, and shopping. CulinaryAI treats the pantry as the source of truth and connects every downstream feature to it."
    )
    document.add_heading("2.1 Primary objectives", level=2)
    add_bullets(
        document,
        [
            "Maintain a household-scoped, real-time inventory that works online and has an offline fallback.",
            "Turn receipt images into reviewable pantry additions instead of silently trusting OCR output.",
            "Recommend verified recipes using available ingredients and explicitly show match and missing-item information.",
            "Guide users through cooking with ordered instructions and ingredient deduction only after confirmation.",
            "Calculate recipe nutrition and scale quantities deterministically for different serving counts.",
            "Provide a conversational entry point for pantry and recipe questions.",
            "Create an AI/ML architecture that can progress from explainable baselines to semantic RAG and preference learning.",
        ],
    )
    document.add_heading("2.2 Team responsibility mapping", level=2)
    add_table(
        document,
        ["Owner", "Proposal responsibility", "Implemented Week 1 result"],
        [
            ["Person A", "App shell, authentication, shared pantry, receipt scanner, shopping links", "Flutter architecture, persistent authentication, household create/join, pantry CRUD/sync, scanner review flow, restock links and guided-cooking deduction hook"],
            ["Person B", "Nutrition data, conversions, recipe arithmetic and portion scaling", "Nutrition models, food reference data, unit conversion, total/per-serving macro calculation, ingredient scaling and UI panel"],
            ["Person C", "FastAPI, pantry-aware planning, guided instructions and RAG embeddings", "FastAPI service, validated schemas, verified recipe dataset, local embedding repository, pantry-aware ranking, generated instruction responses and tests"],
        ],
        [1.0, 2.45, 3.45],
    )

    document.add_heading("3. Implemented System Architecture", level=1)
    add_code_flow(
        document,
        "USER\n  │\n  ▼\nFLUTTER MOBILE CLIENT\n  ├── Authentication & persistent session\n  ├── Household and pantry state (Provider)\n  ├── Receipt review and pantry confirmation\n  ├── Recipe, nutrition and guided cooking UI\n  └── Chef Chat\n        │\n        ├──────────► FIREBASE\n        │             ├── Authentication\n        │             ├── Firestore users / households / pantry\n        │             └── Household-isolation security rules\n        │\n        └──────────► FASTAPI AI SERVICE\n                      ├── /health\n                      ├── /generate-plan\n                      ├── /generate-instructions\n                      └── Verified recipe retrieval pipeline",
    )
    document.add_paragraph(
        "The architecture separates user experience, durable cloud state, deterministic calculations, and AI retrieval. This separation matters because nutrition arithmetic and access control should remain predictable, while retrieval and personalization can evolve independently as AI/ML components improve."
    )
    document.add_heading("3.1 Flutter state flow", level=2)
    document.add_paragraph(
        "The root MaterialApp initializes Firebase and registers providers. AuthProvider determines whether the app shows the loading screen, authentication screens, or MainShell. HouseholdProvider observes the authenticated user, resolves a household or personal pantry scope, and supplies that scope to PantryProvider. PantryProvider then loads cached data and synchronizes remote Firestore updates. This dependency chain prevents pantry data from leaking between accounts or households."
    )
    add_code_flow(
        document,
        "AuthProvider.currentUser\n        ↓\nHouseholdProvider.bindUser(user)\n        ↓\nHouseholdProvider.pantryScopeId\n        ↓\nPantryProvider.bindHousehold(scopeId)\n        ↓\nLocal cache + Firestore household pantry",
    )

    document.add_heading("4. Detailed Work Completed", level=1)
    document.add_heading("4.1 Application shell and visual system", level=2)
    document.add_paragraph(
        "A cohesive Flutter shell now hosts Home, Pantry, Chef Chat, Restock, and Profile. The UI supports light and dark modes using a warm culinary design system. The dashboard contains a personalized chef identity, live pantry count, compact question field, recipe inspiration, storage-zone cards, a receipt-scanning entry point, conditionally visible low-stock content, and pantry-aware recommended meals."
    )
    add_bullets(
        document,
        [
            "The bottom navigation was reduced to a compact 60-pixel height with 20-pixel icons and 10-pixel labels.",
            "Storage-zone badges now use real pantry counts and display zero rather than hardcoded demo values.",
            "The low-stock dashboard block—including its heading, carousel and spacing—is removed when no items are low.",
            "Dashboard and chatbot inputs were reduced to compact single-line controls.",
            "Drawer and profile surfaces use practical navigation rather than decorative AI-themed cards.",
        ],
    )

    document.add_heading("4.2 Authentication and persistent login", level=2)
    document.add_paragraph(
        "Authentication supports email/password registration and login, Google sign-in when Firebase is available, validation, readable error handling, logout confirmation, and profile synchronization. A SessionService stores the signed-in user locally, allowing the app to restore a valid application session on restart. Firebase authentication state is preferred when available; the local session also supports development and offline demonstrations."
    )
    add_numbered(
        document,
        [
            "On startup, Firebase is initialized and AuthProvider checks Firebase Auth or the saved local session.",
            "While restoration is running, a loading screen prevents an incorrect flash of the login page.",
            "If a user is restored, the full profile is loaded and the app opens MainShell.",
            "When the user logs out, Firebase and local session data are both cleared.",
        ],
    )

    document.add_heading("4.3 Personalized chef identity", level=2)
    document.add_paragraph(
        "The dashboard and drawer use a shared ChefNameFormatter. It first uses the saved profile name and takes its first word. When only an email-derived display name is available, it cleans separators and common combined family-name suffixes. For the requested example, waseeqsaad@gmail.com is displayed as Chef Waseeq. The same formatter is used inside Chef Chat so identity remains consistent."
    )

    document.add_heading("4.4 Household creation and shared pantry", level=2)
    document.add_paragraph(
        "A household represents a shared kitchen. A user can create a household and receive a six-character invitation code, or join an existing household with a valid code. Household membership determines which pantry collection the user can read and update. A single user may also use the same concept as a personal kitchen."
    )
    add_bullets(
        document,
        [
            "Household records include owner, members, invitation code, display name and creation time.",
            "The service writes a separate invite record to resolve join requests without exposing all invites.",
            "Offline household data is cached for resilient demos and reconnect behavior.",
            "Firestore rules enforce authenticated access and household membership for pantry operations.",
            "User profile data stores the selected household association.",
        ],
    )

    document.add_heading("4.5 Pantry inventory management", level=2)
    document.add_paragraph(
        "PantryProvider manages the application’s central inventory state. Items include quantity, unit, storage location, low-stock threshold, creation source, update time, household scope and optional custom image. The UI supports adding, editing, deleting, searching, storage-zone filtering and quantity adjustment."
    )
    add_table(
        document,
        ["Capability", "Implementation behavior"],
        [
            ["Storage zones", "Fridge, Freezer, Pantry Shelf and Spice Rack; cards show live counts"],
            ["Search", "Case-insensitive filtering of current pantry item names"],
            ["Low stock", "Quantity is compared with each item’s configurable threshold"],
            ["Offline support", "Local household pantry copy is loaded before/alongside remote sync"],
            ["Remote sync", "Firestore household pantry stream updates Provider state"],
            ["Cooking deduction", "Recipe quantities are deducted only after explicit completion confirmation"],
            ["Receipt source", "Confirmed scanned items are tagged as receipt_scan"],
        ],
        [1.65, 5.25],
    )

    document.add_heading("4.6 Receipt scanning and verified onboarding", level=2)
    document.add_paragraph(
        "The scanner supports camera or gallery image selection, scan/processing visual states, an editable review interface, item selection, quantity and storage-zone confirmation, and batch addition to the pantry. This review-first design is important because OCR can misread product names or quantities and should not modify inventory without user confirmation."
    )
    add_callout(
        document,
        "Current limitation: image capture is real, but Week 1 extraction uses a mock OCR response fixture. The captured image and UI contract are ready for a future vision/OCR endpoint.",
        PALE_ORANGE,
    )

    document.add_heading("4.7 Recipe recommendation and guided cooking", level=2)
    document.add_paragraph(
        "The Flutter application contains a structured local recipe catalog with ingredients, instructions, meal type, time, calories, difficulty, image, and serving information. Recipes calculate their match percentage against the active pantry and report available and missing ingredients. The dashboard selects a hero recipe using pantry match and deterministic daily rotation, while additional recommendations are ordered by pantry coverage."
    )
    document.add_paragraph(
        "Opening a recipe displays detail and nutrition information. Starting a recipe launches GuidedCookingScreen, which progresses through ordered instructions. At completion, the app asks for confirmation before deducting matching ingredients. This avoids accidental inventory changes when a user merely reads or abandons a recipe."
    )

    document.add_heading("4.8 Nutrition and portion scaling", level=2)
    document.add_paragraph(
        "The nutrition module is deliberately deterministic. It uses a reference nutrient database, converts ingredient quantities to grams, accumulates macronutrients, and calculates totals and per-serving values. Ingredient quantities scale by the ratio between requested and base servings. ML is not used here because arithmetic must remain transparent and reproducible."
    )
    add_code_flow(
        document,
        "ingredient quantity + unit\n        ↓ unit normalization\nquantity in grams\n        ↓ × nutrient reference per 100 g\ningredient macros\n        ↓ sum all ingredients\nrecipe total\n        ↓ ÷ servings\nper-serving nutrition",
    )
    add_bullets(
        document,
        [
            "NutritionFacts stores calories, protein, carbohydrates and fat.",
            "UnitConversionService supports mass, volume and count normalization.",
            "NutritionCalculatorService rejects non-positive serving counts.",
            "RecipeNutritionPanel presents total and per-serving results.",
            "Portion scaling updates all structured ingredient quantities consistently.",
        ],
    )

    document.add_heading("4.9 Chef Chat and dashboard question entry", level=2)
    document.add_paragraph(
        "The previous decorative search surface was replaced with an editable dashboard question field. Submitting a question opens Chef Chat and immediately processes that question. The former Recipes placeholder tab was converted into a persistent Chef Chat destination. Conversation state, suggestion chips, message bubbles, typing feedback and clearing are implemented."
    )
    add_table(
        document,
        ["Supported question category", "Example", "Answer source"],
        [
            ["Inventory", "What do I have in my pantry?", "Current PantryProvider items"],
            ["Ingredient lookup", "Do I have milk?", "Quantity, unit and storage location"],
            ["Low stock", "What is running out?", "Item low-stock thresholds"],
            ["Recommendations", "What can I cook?", "Pantry-ranked local recipe catalog"],
            ["Missing ingredients", "What is missing for tomato pasta?", "Recipe-versus-pantry comparison"],
            ["Cooking time", "How long does tomato pasta take?", "Structured recipe metadata"],
            ["Nutrition", "How many calories does it have?", "Recipe calorie metadata"],
            ["Instructions", "How do I cook tomato pasta?", "Ordered recipe instructions"],
        ],
        [1.55, 2.45, 2.9],
    )
    add_callout(
        document,
        "Current limitation: Chef Chat is an offline domain-intent assistant, not an unrestricted LLM conversation. This makes it reliable for known pantry and recipe questions while the production LLM/RAG connection is developed.",
        PALE_ORANGE,
    )

    document.add_heading("4.10 Low-stock restocking", level=2)
    document.add_paragraph(
        "The Restock tab reads live low-stock items. Each item can open an external marketplace search through Daraz or Pandamart using the ingredient name as the query. When all items are sufficiently stocked, the screen displays a clear stocked state instead of an empty list."
    )

    document.add_heading("4.11 FastAPI backend and data contracts", level=2)
    document.add_paragraph(
        "The Python backend is a separate FastAPI service with Pydantic validation, CORS configuration, an indexed verified recipe repository, meal planning, instruction generation, and automatic OpenAPI documentation. It can run locally without a paid API key."
    )
    add_table(
        document,
        ["Endpoint", "Input", "Output"],
        [
            ["GET /health", "No body", "Service status and indexed recipe count"],
            ["POST /generate-plan", "Pantry items, servings, meal count, dietary preferences and optional maximum time", "Ranked recipes, match percentage, available/missing ingredients, scaled quantities and steps"],
            ["POST /generate-instructions", "Recipe ID and servings", "Recipe name, scaled ingredients, ordered steps, timers and completion prompt"],
        ],
        [1.6, 2.8, 2.5],
    )
    document.add_paragraph(
        "Validation constrains servings, meal count, cooking time, ingredient quantities and required text. Unknown recipe IDs return a proper HTTP 404 response. This stable API contract allows the Flutter application to integrate later without depending on the internal retrieval implementation."
    )

    document.add_heading("4.12 Current RAG-style recipe retrieval", level=2)
    document.add_paragraph(
        "The verified recipe knowledge base currently contains six structured recipes. At startup, the repository converts recipe names, cuisines, ingredients, tags and dietary attributes into searchable document text. HashEmbeddingService produces deterministic normalized vectors, and cosine similarity ranks documents against a pantry/preference query. MealPlannerService then combines retrieval relevance with actual pantry coverage."
    )
    add_code_flow(
        document,
        "verified_recipes.json\n        ↓ document text\nlocal deterministic embeddings\n        ↓ in-memory index\nuser pantry + preferences → query embedding\n        ↓ cosine similarity\nretrieved recipes\n        ↓ dietary/time filters\n80% pantry coverage + 20% retrieval score\n        ↓\nranked verified meal plan",
    )
    add_callout(
        document,
        "This is the Week 1 retrieval baseline. Hash embeddings match words reproducibly but do not provide deep semantic understanding. They should not be described as a trained neural embedding model.",
        PALE_ORANGE,
    )

    document.add_heading("4.13 Firebase security and configuration", level=2)
    document.add_paragraph(
        "The repository includes Firebase initialization options, Android Google Services configuration, Firestore deployment configuration and security rules. Rules require authentication, restrict user profiles to their owner, scope pantry access to household members, validate household creation ownership, and allow household joining only through a matching invite record."
    )
    add_bullets(
        document,
        [
            "Users cannot read or update another user’s profile.",
            "Pantry documents are nested under households and require membership.",
            "Invitation collection listing is denied; signed-in clients may resolve a specific code.",
            "Household creation requires the creator to be both owner and member.",
            "Membership updates are constrained to the invite flow or an existing member.",
        ],
    )

    document.add_heading("5. Verification and Quality Evidence", level=1)
    document.add_paragraph(
        "Verification was performed at the service, state-management and application levels. The final Flutter run completed 21 tests successfully. The FastAPI backend completed 8 tests successfully. Static Flutter analysis completed without compilation errors; remaining analyzer messages are non-blocking style/deprecation notices in existing code. Python source compilation and verified recipe JSON parsing also passed."
    )
    add_table(
        document,
        ["Test area", "What is verified", "Result"],
        [
            ["Authentication validation", "Email, password strength, confirmation and phone validation", "Passed"],
            ["Session persistence", "User survives a service restart and is removed on clear", "Passed"],
            ["Households", "Create and join through invitation code", "Passed"],
            ["Pantry", "Offline persistence, provider behavior and controlled recipe deduction", "Passed"],
            ["Nutrition", "Unit normalization, totals, per-serving macros, scaling and invalid input", "Passed"],
            ["Chef naming", "Profile first name, email fallback and waseeqsaad example", "Passed"],
            ["Chef Chat", "Inventory, recommendation, instruction and low-stock answers", "Passed"],
            ["Flutter smoke test", "Application starts at authentication screen", "Passed"],
            ["FastAPI endpoints", "Health, plan response and missing-recipe 404", "Passed"],
            ["Retrieval/planner", "Determinism, relevant retrieval, filters and scaling", "Passed"],
        ],
        [1.45, 4.45, 1.0],
    )
    document.add_paragraph(
        "Verified totals: 21 Flutter tests passed and 8 backend tests passed. Commit c21f739 on main captures the implementation described in this report before the report itself is added."
    )

    document.add_heading("6. Current Limitations", level=1)
    add_table(
        document,
        ["Limitation", "Impact", "Planned resolution"],
        [
            ["Hash-based recipe embeddings", "Weak synonym and semantic understanding", "Use a sentence-transformer model and persistent vector index"],
            ["Small six-recipe backend dataset", "Limited cuisine and dietary coverage", "Curate and validate a substantially larger recipe corpus"],
            ["Mock OCR extraction fixture", "Captured receipt is not yet interpreted by a real model", "Connect vision/OCR endpoint and confidence-aware parser"],
            ["Local intent chatbot", "Cannot answer unrestricted or multi-turn culinary questions", "Connect grounded LLM chat to retrieved pantry/recipe context"],
            ["Flutter/backend not integrated", "Mobile recommendations still use local catalog", "Add HTTP client, environment configuration and error/offline states"],
            ["No learned personalization", "All users receive the same ranking for the same pantry", "Collect feedback and train/evaluate a ranking model"],
            ["Backend not hosted", "FastAPI must run locally", "Deploy to an appropriate Python hosting service"],
            ["Limited automated UI coverage", "Core services are tested more deeply than full interaction flows", "Add widget, integration and Firebase emulator tests"],
        ],
        [1.85, 2.45, 2.6],
    )

    document.add_heading("7. Future AI/ML Architecture", level=1)
    document.add_paragraph(
        "Because this is an AI/ML course project, the next phase should replace placeholder retrieval components with measurable ML while preserving deterministic safety and arithmetic. The recommended design is a hybrid: semantic RAG identifies suitable verified recipes, hard rules enforce constraints, a learned ranker predicts preference, and an optional contextual bandit safely improves ordering from feedback."
    )
    add_code_flow(
        document,
        "pantry + user question + preferences\n        ↓\nneural embedding model\n        ↓\nvector search over verified recipes\n        ↓\nallergy / dietary / availability safety rules\n        ↓\nML preference ranker\n        ↓\ncontextual-bandit re-ranking (safe exploration only)\n        ↓\nrecommendations shown in Flutter\n        ↓\nselect / dismiss / cook / complete / rate\n        └──────────── feedback for future learning",
    )

    document.add_heading("7.1 Semantic RAG upgrade", level=2)
    add_numbered(
        document,
        [
            "Expand and clean the verified recipe corpus, preserving source and structured ingredients.",
            "Generate neural embeddings with a sentence-transformer model such as all-MiniLM-L6-v2 or an equivalent selected through evaluation.",
            "Store vectors in FAISS, Chroma, Qdrant, Pinecone, or Firestore Vector Search depending deployment constraints.",
            "Embed the pantry, free-text request and dietary preferences as the retrieval query.",
            "Retrieve top-k candidates and apply hard dietary/allergy/cooking-time filters.",
            "Provide retrieved context to an LLM only when explanation or conversational composition is needed.",
            "Return citations or recipe IDs so every generated claim is traceable to a verified record.",
        ],
    )
    document.add_paragraph(
        "Semantic retrieval should be evaluated against a labeled query set using Recall@K, Precision@K, Mean Reciprocal Rank and nDCG. Pantry match remains a separate, explainable feature rather than being hidden inside the embedding score."
    )

    document.add_heading("7.2 Supervised recommendation model", level=2)
    document.add_paragraph(
        "After interaction data exists, a supervised ranker can estimate the probability that a user will select and complete a candidate recipe. Suitable initial models include logistic regression, random forest, gradient boosting or a small learning-to-rank model. A simple model is preferable at first because its behavior can be compared against the existing fixed-score baseline."
    )
    add_table(
        document,
        ["Feature group", "Example features"],
        [
            ["Retrieval", "Semantic similarity, retrieval rank"],
            ["Pantry", "Match percentage, missing count, expiry-use score"],
            ["Recipe", "Cuisine, meal type, cook time, difficulty, calories"],
            ["User", "Dietary preference, cuisine history, typical cook duration"],
            ["Interaction", "Past impressions, selections, completions, ratings and dismissals"],
            ["Context", "Time of day, weekday, requested servings and household size"],
        ],
        [1.45, 5.45],
    )

    document.add_heading("7.3 Reinforcement learning through a contextual bandit", level=2)
    document.add_paragraph(
        "Full deep reinforcement learning is unnecessary for this problem. A contextual bandit is a realistic form of reinforcement learning that chooses the order of already safe retrieved recipes. The context is the pantry and user profile, actions are candidate recipes, and rewards come from observable interactions."
    )
    add_table(
        document,
        ["Event", "Illustrative reward", "Reason"],
        [
            ["Recipe displayed", "0", "No evidence of preference"],
            ["Details opened", "+0.5", "Weak interest signal"],
            ["Cook This selected", "+2", "Strong intent"],
            ["Cooking completed", "+5", "Primary success event"],
            ["Positive rating", "+1 to +3", "Explicit satisfaction"],
            ["Immediately dismissed", "−1", "Negative relevance signal"],
            ["Cooking abandoned", "−2", "Recommendation may have been unsuitable"],
        ],
        [2.0, 1.35, 3.55],
    )
    document.add_paragraph(
        "Epsilon-greedy, LinUCB or Thompson Sampling can be compared in a simulated environment before any online use. Exploration must occur only among candidates that already satisfy allergies, dietary restrictions and verified-recipe constraints. The experiment should report cumulative reward, acceptance rate, completion rate and regret against a fixed baseline."
    )

    document.add_heading("7.4 Vision/OCR pipeline", level=2)
    add_code_flow(
        document,
        "receipt image\n   ↓ image preprocessing\nOCR / vision model\n   ↓ raw text + bounding boxes + confidence\nline-item parser\n   ↓ name / quantity / unit / price\ningredient normalization model or dictionary\n   ↓ structured candidates\nuser review and correction\n   ↓\nconfirmed pantry update",
    )
    document.add_paragraph(
        "The review step must remain even after OCR is introduced. User corrections can become labeled examples for later evaluation and improvement, but sensitive receipt content should not be retained longer than required. Metrics should include character error rate, item extraction precision/recall, quantity accuracy, and user correction rate."
    )

    document.add_heading("8. Future Development Roadmap", level=1)
    add_table(
        document,
        ["Phase", "Main work", "Definition of done"],
        [
            ["Phase 2 — integration", "Connect Flutter to FastAPI, add configuration, loading/error/offline states and response mapping", "Mobile app successfully requests and renders backend meal plans and cooking steps"],
            ["Phase 3 — semantic RAG", "Neural embeddings, vector store, larger verified dataset and retrieval evaluation", "Semantic retrieval beats hash baseline on documented Recall@K/nDCG tests"],
            ["Phase 4 — real OCR", "Vision service, structured parser, confidence handling and correction logging", "Real receipts produce editable items with measured extraction accuracy"],
            ["Phase 5 — feedback data", "Record impressions, selections, completions, ratings and dismissals", "Anonymized interaction dataset with documented schema and consent"],
            ["Phase 6 — ML ranking", "Train baseline classifiers/rankers and compare with fixed pantry score", "Selected model improves offline ranking metrics without breaking constraints"],
            ["Phase 7 — contextual bandit", "Safe exploration and reward learning in simulation, then controlled use", "Bandit improves cumulative reward over random/fixed baselines"],
            ["Phase 8 — deployment", "Host FastAPI/vector store, secure environment configuration, monitoring and CI/CD", "Reproducible staging deployment with health monitoring and rollback path"],
            ["Phase 9 — hardening", "Integration tests, accessibility, performance, privacy review and store readiness", "Release candidate meets documented quality gates"],
        ],
        [1.4, 3.45, 2.05],
    )

    document.add_heading("9. Deployment Plan", level=1)
    document.add_heading("9.1 Firebase on the Spark plan", level=2)
    document.add_paragraph(
        "Firebase Authentication and Firestore can support development and a limited course demonstration on the Spark plan, subject to Firebase’s current quotas and enabled-provider requirements. The repository already contains Firebase configuration and Firestore rules. Usage should be monitored, and demo data should remain small."
    )
    document.add_heading("9.2 FastAPI hosting", level=2)
    document.add_paragraph(
        "The FastAPI service cannot run inside Firebase Hosting itself. It requires a Python-capable host such as Cloud Run, Render, Railway, Fly.io, or another container/service platform. Hosting choice should consider free-tier availability, cold starts, memory required by the embedding model, HTTPS, environment-secret handling, logging and budget alerts."
    )
    add_bullets(
        document,
        [
            "Keep API keys and service credentials out of Git; use platform environment variables or secret managers.",
            "Restrict CORS to approved mobile/web origins for production rather than wildcard origins.",
            "Add request timeouts, structured logging, rate limits and model/version reporting.",
            "Use a staging endpoint before production and keep the local deterministic fallback for demonstrations.",
            "Automate Flutter and backend tests before every deployment.",
        ],
    )

    document.add_heading("10. Recommended Evaluation for the AI/ML Course", level=1)
    document.add_paragraph(
        "The project should be presented as an evaluated hybrid system rather than simply an application that calls an AI API. Each intelligent component needs a baseline, dataset, metric and experiment."
    )
    add_table(
        document,
        ["Experiment", "Baseline", "Proposed model", "Metrics"],
        [
            ["Recipe retrieval", "Keyword/hash retrieval", "Sentence-transformer embeddings", "Precision@K, Recall@K, MRR, nDCG"],
            ["Recipe ranking", "80% pantry + 20% retrieval rule", "Logistic regression / boosting ranker", "Top-k accuracy, nDCG, completion prediction AUC"],
            ["Personalization", "Fixed or random ordering", "Epsilon-greedy / Thompson Sampling", "Cumulative reward, regret, acceptance and completion"],
            ["Receipt extraction", "Raw OCR text", "OCR plus structured parser/normalizer", "Item precision/recall, quantity accuracy, correction rate"],
            ["System utility", "Manual recipe browsing", "Full pantry-aware flow", "Task time, successful meal selection and user satisfaction"],
        ],
        [1.35, 1.8, 2.15, 1.6],
    )
    add_bullets(
        document,
        [
            "Split datasets into training, validation and test sets without leaking the same user or near-duplicate recipe across splits.",
            "Report dataset size, class balance, preprocessing and hyperparameters.",
            "Include ablation tests: semantic similarity only, pantry match only, and the combined approach.",
            "Measure both relevance and constraint violations; a useful model must never recommend an allergen because of a high preference score.",
            "Clearly label simulated interaction data and do not present it as real user behavior.",
        ],
    )

    document.add_heading("11. Data, Privacy and Safety Considerations", level=1)
    add_bullets(
        document,
        [
            "Treat dietary restrictions and allergies as hard filters, not probabilistic preferences.",
            "Minimize storage of receipt images and remove personally identifying receipt content where possible.",
            "Use household-scoped authorization for every pantry read and write.",
            "Do not expose invitation collections or sequentially guessable codes through list operations.",
            "Allow users to understand why a recipe was recommended: pantry match, missing items, time and preference factors.",
            "Version models and recipe datasets so results can be reproduced and rolled back.",
            "Collect explicit consent before using interaction histories for model training.",
            "Avoid presenting calculated nutrition as medical advice; reference data and assumptions should be disclosed.",
        ],
    )

    document.add_heading("12. Repository Guide", level=1)
    add_table(
        document,
        ["Path", "Purpose"],
        [
            ["lib/main.dart", "Firebase startup, provider dependency graph and authentication gate"],
            ["lib/screens/home/main_shell.dart", "Dashboard, navigation, recipe surfaces, restock and drawer"],
            ["lib/screens/chat/chatbot_screen.dart", "Chef Chat interface and conversation state"],
            ["lib/services/chef_chat_service.dart", "Local pantry/recipe question-answering logic"],
            ["lib/providers/pantry_provider.dart", "Pantry state, filtering, persistence, synchronization and quantity changes"],
            ["lib/services/household_service.dart", "Household create/join and invitation behavior"],
            ["lib/services/nutrition_calculator_service.dart", "Nutrition totals, per-serving values and ingredient scaling"],
            ["lib/screens/scanner/receipt_scanner_screen.dart", "Capture, processing, review and confirmation workflow"],
            ["backend/app/main.py", "FastAPI application and routes"],
            ["backend/app/services/meal_planner.py", "Filtering, pantry scoring, scaling and instruction response"],
            ["backend/app/services/recipe_repository.py", "Verified recipe loading, indexing and retrieval"],
            ["backend/app/data/verified_recipes.json", "Structured verified recipe knowledge base"],
            ["firestore.rules", "User and household data authorization"],
            ["test/ and backend/tests/", "Flutter and Python automated verification"],
        ],
        [2.65, 4.25],
    )

    document.add_heading("13. How to Run the Current System", level=1)
    document.add_heading("13.1 Flutter application", level=2)
    add_code_flow(document, "flutter pub get\nflutter test\nflutter run")
    document.add_heading("13.2 FastAPI backend", level=2)
    add_code_flow(
        document,
        "python -m venv backend\\.venv\nbackend\\.venv\\Scripts\\python.exe -m pip install -r backend\\requirements.txt\ncd backend\n.\\.venv\\Scripts\\python.exe -m uvicorn app.main:app --reload",
    )
    document.add_paragraph(
        "After startup, interactive backend documentation is available at http://127.0.0.1:8000/docs. Backend tests run with .\\.venv\\Scripts\\python.exe -m pytest -q from the backend directory."
    )

    document.add_heading("14. Definition of the Next Successful Milestone", level=1)
    document.add_paragraph(
        "The next milestone should not be defined merely as adding more screens. It should demonstrate one complete intelligent loop from real application data to measurable AI behavior."
    )
    add_numbered(
        document,
        [
            "Flutter sends the authenticated household pantry and a user question to FastAPI.",
            "FastAPI retrieves recipes using a real semantic embedding model and persistent vector index.",
            "Hard constraints remove unsafe or incompatible recipes.",
            "The response is rendered in Chef Chat and can launch Guided Cooking.",
            "Selection and completion feedback is recorded using a documented event schema.",
            "Retrieval quality is compared against the current hash/keyword baseline using a held-out test set.",
            "A report presents metrics, limitations, examples and reproducible commands.",
        ],
    )
    add_callout(
        document,
        "Recommended next deliverable: semantic RAG integrated end-to-end with Flutter, accompanied by a small labeled retrieval benchmark and comparison against the current baseline.",
    )

    document.add_heading("15. Conclusion", level=1)
    document.add_paragraph(
        "CulinaryAI now has a credible Week 1 foundation across mobile UX, authentication, household data, pantry management, receipt onboarding, nutrition, recipes, guided cooking, conversational access and a separate AI service. The strongest engineering decision is the separation between deterministic operations—authorization, inventory and nutrition—and evolvable AI operations—retrieval, ranking, vision and personalization."
    )
    document.add_paragraph(
        "The project is not yet a finished production AI system, and this report intentionally avoids that claim. Its current baselines are functional, tested and structured so they can be replaced with measurable ML components. Completing the semantic RAG integration, real OCR, feedback dataset, supervised ranking and safe contextual-bandit experiments will turn the existing application into a defensible AI/ML course project with both practical value and experimental depth."
    )

    document.add_page_break()
    document.add_heading("Appendix A — Current API Request Example", level=1)
    add_code_flow(
        document,
        '{\n  "pantry_items": [\n    {"name": "chicken breast", "quantity": 500, "unit": "g"},\n    {"name": "basmati rice", "quantity": 250, "unit": "g"}\n  ],\n  "servings": 4,\n  "meal_count": 3,\n  "dietary_preferences": [],\n  "max_cook_time_minutes": 45\n}',
    )
    document.add_heading("Appendix B — Current Verification Commands", level=1)
    add_code_flow(
        document,
        "flutter test\nflutter analyze --no-fatal-infos --no-fatal-warnings\n\ncd backend\n.\\.venv\\Scripts\\python.exe -m pytest -q\n.\\.venv\\Scripts\\python.exe -m compileall -q app tests",
    )
    return document


def validate_docx(path: Path) -> None:
    if not path.exists() or path.stat().st_size < 20_000:
        raise RuntimeError("Generated DOCX is missing or unexpectedly small")
    with ZipFile(path) as archive:
        required = {"[Content_Types].xml", "word/document.xml", "word/styles.xml"}
        missing = required.difference(archive.namelist())
        if missing:
            raise RuntimeError(f"Generated DOCX misses required parts: {sorted(missing)}")
        document_xml = archive.read("word/document.xml").decode("utf-8")
        for phrase in (
            "Executive Summary",
            "Detailed Work Completed",
            "Future AI/ML Architecture",
            "Future Development Roadmap",
            "Current Limitations",
        ):
            if phrase not in document_xml:
                raise RuntimeError(f"Generated DOCX misses expected section: {phrase}")


if __name__ == "__main__":
    report = build_report()
    report.save(OUTPUT)
    validate_docx(OUTPUT)
    print(f"Generated {OUTPUT} ({OUTPUT.stat().st_size:,} bytes)")
